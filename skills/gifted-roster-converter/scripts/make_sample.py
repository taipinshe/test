#!/usr/bin/env python3
"""產生一份「假的」成績輸入檔，結構仿照市府發給各校的成績輸入表，供測試與示範用。

裡面沒有任何真實學生資料。仿照的重點（都是實際檔案出現過、會讓轉檔出錯的地方）：
- 工作表名稱前後有空格（例如 '數學薦送名單 '、'自然科學薦送名單  '）
- 表頭不在第一列，上面有標題列
- 班級是全形文字（'７０１'），座號有前導零
- 特殊身分欄的空白格是一串空格，不是真正的空值
- 有人缺一次段考分數

用法：python make_sample.py 輸出檔.xlsx [--students 60] [--seed 7]
"""
import argparse
import random

from openpyxl import Workbook

SHEETS = [('國文', '國文PR計算', '國文薦送名單'),
          ('英語', '英語PR計算 ', '英語薦送名單'),
          ('數學', '數學PR計算', '數學薦送名單 '),
          ('自然', '自然科學PR計算', '自然科學薦送名單  ')]
HEADER = ['班級', '座號', '學生姓名', '特殊身分標記', '段考一', '段考二', '段考三', '總分', 'Pr值']
MARKS = {3: '原住民', 11: '中低收入戶', 20: '身心障礙', 33: '新住民'}   # 學生序號 -> 特殊身分


def to_fullwidth(s):
    return ''.join(chr(ord(c) + 0xFEE0) if c.isdigit() else c for c in s)


def percentile_rank(totals):
    n = len(totals)
    return [((sum(1 for t in totals if t < x) * 1000) // (n - 1)) / 10 for x in totals]


def build(n_students=60, seed=7):
    rng = random.Random(seed)
    roster = []
    for i in range(n_students):
        cls = 701 + i // 30
        roster.append({'班級': to_fullwidth(str(cls)), '座號': f'{i % 30 + 1:02d}',
                       '學生姓名': f'測試生{i + 1:03d}', '身分': MARKS.get(i, ' ' * 8)})

    wb = Workbook()
    wb.remove(wb.active)
    for subject_no, (subject, pr_name, rec_name) in enumerate(SHEETS):
        scores = []
        for i in range(n_students):
            row = [rng.randint(40, 100) for _ in range(3)]
            if subject_no == 2 and i == 5:   # 數學有一位缺考一次
                row[1] = None
            scores.append(row)
        totals = [sum(s for s in row if s is not None) for row in scores]
        prs = percentile_rank(totals)

        ws = wb.create_sheet(pr_name)
        ws.append([f'114學年度第一學期 {subject} 定期評量成績（範例資料）'])
        ws.append(HEADER)
        for stu, row, total, pr in zip(roster, scores, totals, prs):
            ws.append([stu['班級'], stu['座號'], stu['學生姓名'], stu['身分'], *row, total, pr])

        rec = wb.create_sheet(rec_name)
        rec.append([f'{subject} 薦送名單（範例資料）'])
        rec.append(HEADER)
        for stu, row, total, pr in zip(roster, scores, totals, prs):
            special = stu['身分'].strip() != ''
            if pr >= 93 or (special and pr >= 85):
                rec.append([stu['班級'], stu['座號'], stu['學生姓名'], stu['身分'], *row, total, pr])
    return wb


def main():
    p = argparse.ArgumentParser(description='產生假的資優成績輸入檔')
    p.add_argument('out')
    p.add_argument('--students', type=int, default=60)
    p.add_argument('--seed', type=int, default=7)
    args = p.parse_args()
    build(args.students, args.seed).save(args.out)
    print(f'已產生 {args.out}')


if __name__ == '__main__':
    main()
