"""轉檔測試。全部用 make_sample.py 產生的假資料，不含真實學生。

執行：python -m pytest skills/gifted-roster-converter/tests -q
"""
import io
import os
import sys

import pytest
from openpyxl import load_workbook

HERE = os.path.dirname(__file__)
sys.path.insert(0, os.path.join(HERE, '..', 'scripts'))

import convert as cv          # noqa: E402
import make_sample            # noqa: E402


def to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run(wb, tmp_path, filename='sample.xlsx', **kw):
    return cv.convert(filename, to_bytes(wb), out_dir=str(tmp_path), **kw)


def test_percentile_rank_matches_percentrank_inc():
    # PERCENTRANK.INC({10,20,20,30}, x)：10→0、20→0.333→33.3、30→1
    assert cv.percentile_rank([10, 20, 20, 30]) == [0.0, 33.3, 33.3, 100.0]
    # 無條件捨去：贏 2 人 ÷ 6 = 0.3333… → 33.3；贏 4 人 ÷ 6 = 0.6666… → 66.6（不是 66.7）
    assert cv.percentile_rank([1, 2, 3, 4, 5, 6, 7])[2] == 33.3
    assert cv.percentile_rank([1, 2, 3, 4, 5, 6, 7])[4] == 66.6


def test_sample_converts_and_matches_school_list(tmp_path):
    report = run(make_sample.build(), tmp_path)
    assert report['合計人數'] > 0
    for s in report['各科']:
        assert s['學校名單多列'] == [] and s['學校名單漏列'] == [], s
        assert s['全年級人數'] == 60
    # 只該提醒缺考那一筆
    assert len(report['請人工確認']) == 1 and '數學' in report['請人工確認'][0]

    ws = load_workbook(report['輸出檔'])['薦送名單']
    rows = list(ws.values)
    assert list(rows[0]) == cv.OUT_COLUMNS
    assert all(r[0] == 115 for r in rows[1:])
    assert all(r[1] is None and r[3] is None for r in rows[1:])     # 學校代碼、學號留空
    assert all(set(r[4:]) <= {1, None} and 1 in r[4:] for r in rows[1:])


def test_uppercase_extension_still_writes_lowercase_output(tmp_path):
    report = run(make_sample.build(), tmp_path, filename='115國中資優成績輸入-某校0127.XLSX')
    assert report['輸出檔'].endswith('已轉檔_115國中資優成績輸入-某校0127.xlsx')
    assert os.path.exists(report['輸出檔'])


def test_missing_subject_refuses_and_writes_nothing(tmp_path):
    wb = make_sample.build()
    del wb['數學PR計算']
    with pytest.raises(cv.ConvertError, match='數學'):
        run(wb, tmp_path)
    assert os.listdir(tmp_path) == []


def test_flags_school_list_and_pr_mistakes(tmp_path):
    wb = make_sample.build()
    rec = wb['國文薦送名單']
    pr_ws = wb['國文PR計算']
    rows = list(pr_ws.iter_rows(min_row=3, values_only=True))
    listed = {r[2] for r in rec.iter_rows(min_row=3, values_only=True)}
    lowest = min(rows, key=lambda r: r[8])
    rec.append(list(lowest))                         # 學校多列一個不合格的人
    for row in rec.iter_rows(min_row=3):
        if row[2].value in listed:
            rec.delete_rows(row[0].row)              # 學校漏列一個合格的人
            break
    pr_ws.cell(row=3, column=9, value=12.3)          # 學校表格的 Pr值 寫錯一格
    pr_ws.cell(row=4, column=4, value='無')          # 身分欄出現簡章沒有的寫法

    problems = '\n'.join(run(wb, tmp_path)['請人工確認'])
    assert '學校名單有、但重算不符合資格' in problems
    assert '重算符合資格、但學校名單沒有列' in problems
    assert 'Pr值」與重算結果不同' in problems
    assert "['無']" in problems


def test_special_status_uses_lower_threshold(tmp_path):
    # 一般門檻設到不可能達到、特殊身分門檻設 0：剩下的一定都是有特殊身分標記的人
    report = run(make_sample.build(), tmp_path, pr_general=101, pr_special=0)
    special_names = {f'測試生{i + 1:03d}' for i in make_sample.MARKS}
    names = {r[2] for r in list(load_workbook(report['輸出檔'])['薦送名單'].values)[1:]}
    assert names == special_names


def test_cli_exit_codes(tmp_path):
    good = tmp_path / 'good.xlsx'
    make_sample.build().save(good)
    bad_wb = make_sample.build()
    del bad_wb['英語PR計算 ']
    bad = tmp_path / 'bad.xlsx'
    bad_wb.save(bad)
    out = tmp_path / 'out'
    assert cv.main([str(good), '--out-dir', str(out)]) == 0
    assert cv.main([str(good), str(bad), '--out-dir', str(out), '--json']) == 1
